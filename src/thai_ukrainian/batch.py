from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable

from .api import parse_thai


def analyze_input(text: str) -> dict[str, Any]:
    analyses = parse_thai(text)
    return {
        "input": text,
        "status": "analyzed" if analyses and all(a.status == "analyzed" for a in analyses) else (
            analyses[0].status if len(analyses) == 1 else "mixed"
        ),
        "analyses": [a.as_dict() for a in analyses],
        "ipa": " ".join(a.phonemic_ipa or "?" for a in analyses),
        "phonetic_ipa": " ".join(a.phonetic_ipa or "?" for a in analyses),
        "tones": [a.tone.tone if a.tone else None for a in analyses],
        "tone_ipa": [a.tone.contour_ipa if a.tone else None for a in analyses],
        "ukrainian_orthography": [
            a.selected_ukrainian_orthography for a in analyses
        ],
        "warnings": [w for a in analyses for w in a.warnings],
    }


def read_txt(path: Path) -> list[dict[str, str]]:
    rows = []
    for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        text = line.strip()
        if text:
            rows.append({"id": str(number), "thai": text})
    return rows


def read_csv(path: Path, column: str | None = None) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError("CSV input has no header row.")
        field = column or ("thai" if "thai" in reader.fieldnames else reader.fieldnames[0])
        if field not in reader.fieldnames:
            raise ValueError(f"CSV column {field!r} was not found.")
        return [
            {"id": str(i), "thai": (row.get(field) or "").strip()}
            for i, row in enumerate(reader, 1)
            if (row.get(field) or "").strip()
        ]


def read_xlsx(path: Path, column: str | None = None, sheet: str | None = None) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel support requires openpyxl. Install with: "
            "pip install 'thai-ukrainian-phonetic-system[excel]'"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(x).strip() if x is not None else "" for x in next(rows)]
    except StopIteration:
        raise ValueError("Excel sheet is empty.")
    field = column or ("thai" if "thai" in headers else next((h for h in headers if h), None))
    if not field or field not in headers:
        raise ValueError(f"Excel column {field!r} was not found.")
    index = headers.index(field)
    result = []
    for i, row in enumerate(rows, 2):
        value = row[index] if index < len(row) else None
        text = str(value).strip() if value is not None else ""
        if text:
            result.append({"id": str(i), "thai": text})
    workbook.close()
    return result


def load_batch(path: str | Path, column: str | None = None, sheet: str | None = None) -> list[dict[str, str]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return read_txt(path)
    if suffix in {".csv", ".tsv"}:
        return read_csv(path, column)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path, column, sheet)
    raise ValueError("Supported batch inputs: .txt, .csv, .tsv, .xlsx, .xlsm")


def analyze_rows(rows: Iterable[dict[str, str]]) -> list[dict[str, Any]]:
    output = []
    for row in rows:
        result = analyze_input(row["thai"])
        result["id"] = row.get("id")
        output.append(result)
    return output


def write_jsonl(results: list[dict[str, Any]], path: Path) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for result in results:
            handle.write(json.dumps(result, ensure_ascii=False, default=str) + "\n")


def write_json(results: list[dict[str, Any]], path: Path) -> None:
    path.write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")


def write_csv(results: list[dict[str, Any]], path: Path) -> None:
    fields = [
        "id", "input", "status", "ipa", "phonetic_ipa",
        "tones", "tone_ipa", "ukrainian_orthography", "warnings",
    ]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for result in results:
            writer.writerow({
                "id": result["id"],
                "input": result["input"],
                "status": result["status"],
                "ipa": result["ipa"],
                "phonetic_ipa": result["phonetic_ipa"],
                "tones": "; ".join(x or "" for x in result["tones"]),
                "tone_ipa": "; ".join(x or "" for x in result["tone_ipa"]),
                "ukrainian_orthography": "; ".join(x or "" for x in result["ukrainian_orthography"]),
                "warnings": " | ".join(result["warnings"]),
            })


def write_xlsx(results: list[dict[str, Any]], path: Path) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel support requires openpyxl. Install with: "
            "pip install 'thai-ukrainian-phonetic-system[excel]'"
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Thai-UA"
    fields = [
        "id", "input", "status", "ipa", "phonetic_ipa",
        "tones", "tone_ipa", "ukrainian_orthography", "warnings",
    ]
    sheet.append(fields)
    for result in results:
        sheet.append([
            result["id"], result["input"], result["status"], result["ipa"],
            result["phonetic_ipa"],
            "; ".join(x or "" for x in result["tones"]),
            "; ".join(x or "" for x in result["tone_ipa"]),
            "; ".join(x or "" for x in result["ukrainian_orthography"]),
            " | ".join(result["warnings"]),
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 60)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)


def write_batch(results: list[dict[str, Any]], path: str | Path) -> None:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        write_jsonl(results, path)
    elif suffix == ".json":
        write_json(results, path)
    elif suffix == ".csv":
        write_csv(results, path)
    elif suffix == ".xlsx":
        write_xlsx(results, path)
    else:
        raise ValueError("Supported batch outputs: .json, .jsonl, .csv, .xlsx")
